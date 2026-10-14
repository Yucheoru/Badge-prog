"""L-café オープンバッジの授与対象者を判定するツール。

参加記録の名簿から学生×プログラムごとの活動時間を集計し、
Gold / Silver / Bronze の対象者を Excel に出力する。
"""
import re
import sys
from collections import Counter

import pandas as pd
from openpyxl.styles import PatternFill

PARTICIPATION_SHEET = 'participation_ALL_clean'
REGISTRATION_SHEET = 'OpenBadge_all_clean'

# プログラム名にキーワードがすべて含まれる最初のルールを使う。
# 名簿のプログラム名には "L-café" などの接頭辞が付くため部分一致にしている。
BADGE_CRITERIA_RULES = [
    (['jcpp', 'learner'], {'Gold': 30, 'Silver': 20, 'Bronze': 10}),
    (['jcpp', 'provider'], {'Gold': 45, 'Silver': 30, 'Bronze': 15}),
]
DEFAULT_CRITERIA = {'Gold': 45, 'Silver': 30, 'Bronze': 15}

BADGE_FILLS = {
    'Gold': PatternFill(fill_type='solid', fgColor='FFD700'),
    'Silver': PatternFill(fill_type='solid', fgColor='C0C0C0'),
    'Bronze': PatternFill(fill_type='solid', fgColor='CD7F32'),
}
HALF_HOUR_FILL = PatternFill(fill_type='solid', fgColor='CCE5FF')


def normalize_text(value) -> str:
    """表記ゆれを吸収するため、空白（全角含む）と括弧まわりのスペースを揃える。"""
    s = re.sub(r'[\s　]+', ' ', str(value)).strip()
    s = re.sub(r'\s+（', '（', s)
    s = re.sub(r'）\s+', '）', s)
    return s


def normalize_sid(value) -> str:
    """学生番号は大文字・小文字が混在するので大文字に揃える。"""
    return normalize_text(value).upper()


def find_column(df: pd.DataFrame, *keywords: str):
    """列名にキーワードがすべて含まれる最初の列を返す。なければ None。"""
    keywords = [k.lower() for k in keywords]
    for col in df.columns:
        if all(k in col.lower() for k in keywords):
            return col
    return None


def determine_badge(hours: float, program: str) -> str:
    program = str(program).lower()
    criteria = DEFAULT_CRITERIA
    for keywords, rule in BADGE_CRITERIA_RULES:
        if all(k in program for k in keywords):
            criteria = rule
            break
    for badge in ('Gold', 'Silver', 'Bronze'):
        if hours >= criteria[badge]:
            return badge
    return ''


def load_registrations(xl: pd.ExcelFile, log):
    """バッジ登録シートから (学生番号, プログラム名) ごとの登録日・メール・卒業フラグを読む。

    同じ学生が複数回登録している場合は最も古い登録日を使う。
    """
    reg_dates, emails, graduated = {}, {}, set()

    if REGISTRATION_SHEET not in xl.sheet_names:
        log(f'警告: {REGISTRATION_SHEET} シートがないため、登録日とメールは空欄になります。')
        log(f'シート一覧: {xl.sheet_names}')
        return reg_dates, emails, graduated

    df = xl.parse(REGISTRATION_SHEET)
    col_sid = find_column(df, '学生番号') or find_column(df, 'student')
    col_prog = find_column(df, 'プログラム')
    col_date = find_column(df, 'オープンバッジ登録日') or find_column(df, '登録日')
    col_email = find_column(df, 'email') or find_column(df, 'メール')
    col_grad = find_column(df, '卒業') or find_column(df, 'gold')

    if not (col_sid and col_prog and col_date):
        log(f'警告: {REGISTRATION_SHEET} の必要な列が見つかりません。列: {list(df.columns)}')
        return reg_dates, emails, graduated

    df[col_date] = pd.to_datetime(df[col_date], errors='coerce')
    for _, row in df.iterrows():
        key = (normalize_sid(row[col_sid]), normalize_text(row[col_prog]))
        date = row[col_date]
        if pd.notna(date) and (key not in reg_dates or date < reg_dates[key]):
            reg_dates[key] = date
        if col_email and pd.notna(row.get(col_email)) and key not in emails:
            emails[key] = row[col_email]
        if col_grad and pd.to_numeric(row.get(col_grad), errors='coerce') == 1:
            graduated.add(key)

    log(f'登録日を {len(reg_dates)} 件取得しました（学生番号とプログラム名で照合）。')
    if col_grad:
        log(f'卒業/Gold の {len(graduated)} 件は別シートに出力します。')
    else:
        log('卒業/Gold 列がないため、別シートへの分離は行いません。')
    return reg_dates, emails, graduated


def write_result_sheet(writer, rows: list, sheet_name: str) -> None:
    """結果を書き出し、バッジ列と 0.5 時間の参加日セルに色を付ける。"""
    half_hour_days = [row.pop('_half_hour_days') for row in rows]
    df = pd.DataFrame(rows)
    df.to_excel(writer, sheet_name=sheet_name, index=False)
    if df.empty:
        return

    ws = writer.sheets[sheet_name]
    badge_col = df.columns.get_loc('取得バッジ') + 1
    for r in range(2, len(rows) + 2):
        cell = ws.cell(row=r, column=badge_col)
        if cell.value in BADGE_FILLS:
            cell.fill = BADGE_FILLS[cell.value]

    date_cols = {int(c[3:]): df.columns.get_loc(c) + 1 for c in df.columns if c.startswith('参加日')}
    for i, days in enumerate(half_hour_days):
        for j in days:
            ws.cell(row=i + 2, column=date_cols[j]).fill = HALF_HOUR_FILL


def process(input_path: str, output_path: str, log=print) -> None:
    log('読み込み中...')
    xl = pd.ExcelFile(input_path)
    if PARTICIPATION_SHEET not in xl.sheet_names:
        raise ValueError(f'{PARTICIPATION_SHEET} シートが見つかりません。シート一覧: {xl.sheet_names}')
    df = xl.parse(PARTICIPATION_SHEET)
    total_input = len(df)
    log(f'参加記録: {total_input} 行')

    reg_dates, emails, graduated = load_registrations(xl, log)

    col_sid = find_column(df, '学生番号') or find_column(df, '学籍番号')
    col_name = find_column(df, '氏名') or find_column(df, 'name')
    col_date = find_column(df, '参加日')
    col_hours = find_column(df, '1のみ') or find_column(df, '結果統合', '計算')
    col_program = find_column(df, 'badge_status') or find_column(df, 'badge', 'status')
    col_reason = find_column(df, 'exclude_reason') or find_column(df, '除外')

    missing = [label for label, col in [
        ('学生番号', col_sid), ('氏名', col_name), ('参加日', col_date),
        ('結果統合', col_hours), ('badge_status', col_program),
    ] if col is None]
    if missing:
        raise ValueError(f'必要な列が見つかりません: {missing}\n列一覧: {list(df.columns)}')

    # 結果統合列は 1（参加）/ 0.5（半分参加）/ 0 または空欄（対象外）
    log(f'結果統合列の値: {df[col_hours].value_counts(dropna=False).to_dict()}')
    df[col_hours] = pd.to_numeric(df[col_hours], errors='coerce')
    valid_mask = df[col_hours] > 0
    df_excluded = df[~valid_mask].copy()
    df_valid = df[valid_mask].copy()
    log(f'集計対象: {len(df_valid)} 行 / 対象外（0 以下・空欄）: {len(df_excluded)} 行')

    blank_sid = df_valid[col_sid].isna() | df_valid[col_sid].astype(str).str.strip().isin(['', 'nan', 'None'])
    if blank_sid.any():
        log(f'注意: 学生番号が空欄の行が {blank_sid.sum()} 行あります（1 人分としてまとめて集計されます）。')

    df_valid[col_sid] = df_valid[col_sid].astype(str).str.strip()
    df_valid[col_date] = pd.to_datetime(df_valid[col_date], errors='coerce', format='mixed')
    unreadable_dates = int(df_valid[col_date].isna().sum())
    if unreadable_dates:
        log(f'注意: 参加日を読み取れない行が {unreadable_dates} 行あり、集計から外しました。')

    # 学生番号とプログラム名の組み合わせごとに参加記録をまとめる
    records = {}
    for _, row in df_valid.iterrows():
        key = (normalize_sid(row[col_sid]), normalize_text(row[col_program]))
        rec = records.setdefault(key, {'names': Counter(), 'entries': []})
        rec['names'][row[col_name]] += 1
        if pd.notna(row[col_date]):
            rec['entries'].append((row[col_date], row[col_hours]))
    log(f'{len(records)} 件（学生×プログラム）を判定します。')

    rows, graduated_rows = [], []
    removed_before_reg = 0
    counted_days = 0
    for (sid, program), rec in records.items():
        reg_date = reg_dates.get((sid, program))
        entries = sorted(rec['entries'], key=lambda e: e[0])
        if reg_date is not None:
            # バッジ登録日より前の参加は数えない
            kept = [(d, h) for d, h in entries if d >= reg_date]
            removed_before_reg += len(entries) - len(kept)
            entries = kept
        counted_days += len(entries)
        total_hours = sum(h for _, h in entries)

        row = {
            'No': '',
            'プログラム名': program,
            '学部': '',
            '学生番号': sid,
            '学生氏名（学生申請）': rec['names'].most_common(1)[0][0],
            '学生氏名（英文）': '',
            '学生氏名': '',
            'メールアドレス（学生申請）': emails.get((sid, program), ''),
            'メールアドレス（岡大公式）': '',
            '年次': '',
            '国籍': '',
            '備考': '',
            '言語': '',
            '取得バッジ': determine_badge(total_hours, program),
            'オープンバッジ登録日': reg_date if reg_date is not None else '',
            '合計活動時間': total_hours,
        }
        for j, (d, _) in enumerate(entries, start=1):
            row[f'参加日{j}'] = d.strftime('%Y-%m-%d')
        row['_half_hour_days'] = {j for j, (_, h) in enumerate(entries, start=1) if h == 0.5}

        (graduated_rows if (sid, program) in graduated else rows).append(row)

    for group in (rows, graduated_rows):
        group.sort(key=lambda r: str(r['プログラム名']))
        for i, r in enumerate(group, start=1):
            r['No'] = i

    # 対象外の行も確認できるように別シートへ残す
    df_excluded['メールアドレス'] = [
        emails.get((normalize_sid(s), normalize_text(p)), '')
        for s, p in zip(df_excluded[col_sid], df_excluded[col_program])
    ]
    df_excluded['No'] = range(1, len(df_excluded) + 1)
    excluded_cols = {
        'No': 'No', col_program: 'プログラム名', col_sid: '学生番号', col_name: '氏名',
        'メールアドレス': 'メールアドレス', col_date: '参加日', col_hours: '結果統合値',
    }
    if col_reason:
        excluded_cols[col_reason] = '除外理由'
    df_excluded_out = df_excluded[list(excluded_cols)].rename(columns=excluded_cols).reset_index(drop=True)

    # 入力行がすべてどこかに振り分けられたかを確認する
    accounted = len(df_excluded) + unreadable_dates + removed_before_reg + counted_days
    log('[件数チェック]')
    log(f'入力行数: {total_input}')
    log(f'  対象外（結果統合が 0 以下・空欄）: {len(df_excluded)}')
    log(f'  参加日を読み取れず除外: {unreadable_dates}')
    log(f'  バッジ登録日より前のため除外: {removed_before_reg}')
    log(f'  集計した参加日数: {counted_days}')
    log(f'検算: {accounted} / {total_input}（{"OK" if accounted == total_input else "不一致"}）')

    with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
        write_result_sheet(writer, rows, '結果')
        write_result_sheet(writer, graduated_rows, '卒業_Gold')
        df_excluded_out.to_excel(writer, sheet_name='除外データ', index=False)

    log(f'完了: 結果 {len(rows)} 人 / 卒業_Gold {len(graduated_rows)} 人 → {output_path}')
    log(f'対象外の {len(df_excluded_out)} 行は「除外データ」シートに出力しました。')


if __name__ == '__main__':
    if len(sys.argv) != 3:
        print('使い方: python main.py <入力.xlsx> <出力.xlsx>')
        sys.exit(1)
    process(sys.argv[1], sys.argv[2])
