from datetime import date

import pandas as pd
import pytest
from openpyxl import load_workbook

from main import determine_badge, normalize_sid, normalize_text, process


@pytest.mark.parametrize('hours, program, expected', [
    (9.5, 'L-café JCPP Learner', ''),
    (10, 'L-café JCPP Learner', 'Bronze'),
    (20, 'L-café JCPP Learner', 'Silver'),
    (30, 'L-café JCPP Learner', 'Gold'),
    (29.5, 'JCPP Provider', 'Bronze'),
    (30, 'JCPP Provider', 'Silver'),
    (45, 'jcpp provider', 'Gold'),
    (14.5, 'Language Exchange', ''),
    (15, 'Language Exchange', 'Bronze'),
])
def test_determine_badge(hours, program, expected):
    assert determine_badge(hours, program) == expected


def test_normalize_text_absorbs_spacing():
    assert normalize_text('L-café　JCPP  Learner ') == 'L-café JCPP Learner'
    assert normalize_text('JCPP （Learner） A') == 'JCPP（Learner）A'
    assert normalize_sid(' s1001 ') == 'S1001'


def make_input(path):
    participation = pd.DataFrame([
        # S1: 12 時間参加、うち 2 時間はバッジ登録日より前
        *[{'学生番号': 's1', '氏名': 'Aiko', '参加日': date(2026, 4, d), '結果統合（1のみ計算）': 1,
           'badge_status': 'L-café JCPP Learner', 'exclude_reason': ''} for d in range(1, 13)],
        {'学生番号': 'S1', '氏名': 'Aiko', '参加日': date(2026, 4, 20), '結果統合（1のみ計算）': 0.5,
         'badge_status': 'L-café JCPP Learner', 'exclude_reason': ''},
        # 欠席扱いの行は除外データへ
        {'学生番号': 's1', '氏名': 'Aiko', '参加日': date(2026, 4, 21), '結果統合（1のみ計算）': 0,
         'badge_status': 'L-café JCPP Learner', 'exclude_reason': '欠席'},
        # S2: 卒業/Gold フラグ付き
        {'学生番号': 's2', '氏名': 'Ben', '参加日': date(2026, 5, 1), '結果統合（1のみ計算）': 1,
         'badge_status': 'Language Exchange', 'exclude_reason': ''},
    ])
    registration = pd.DataFrame([
        {'学生番号': 'S1', 'プログラム名': 'L-café JCPP Learner', 'オープンバッジ登録日': date(2026, 4, 3),
         'email': 'aiko@example.com', '卒業/Gold': None},
        {'学生番号': 's2', 'プログラム名': 'Language Exchange', 'オープンバッジ登録日': date(2026, 4, 1),
         'email': 'ben@example.com', '卒業/Gold': 1},
    ])
    with pd.ExcelWriter(path) as w:
        participation.to_excel(w, sheet_name='participation_ALL_clean', index=False)
        registration.to_excel(w, sheet_name='OpenBadge_all_clean', index=False)


def test_process_end_to_end(tmp_path):
    src, dst = tmp_path / 'in.xlsx', tmp_path / 'out.xlsx'
    make_input(src)
    logs = []
    process(str(src), str(dst), log=logs.append)

    result = pd.read_excel(dst, sheet_name='結果')
    assert len(result) == 1
    s1 = result.iloc[0]
    assert s1['学生番号'] == 'S1'
    assert s1['合計活動時間'] == 10.5  # 4/1, 4/2 は登録日前なので除外
    assert s1['取得バッジ'] == 'Bronze'
    assert s1['メールアドレス（学生申請）'] == 'aiko@example.com'

    graduated = pd.read_excel(dst, sheet_name='卒業_Gold')
    assert list(graduated['学生番号']) == ['S2']

    excluded = pd.read_excel(dst, sheet_name='除外データ')
    assert list(excluded['除外理由']) == ['欠席']

    assert any(line.startswith('検算') and 'OK' in line for line in logs)

    ws = load_workbook(dst)['結果']
    header = [c.value for c in ws[1]]
    badge_cell = ws.cell(row=2, column=header.index('取得バッジ') + 1)
    assert badge_cell.fill.fgColor.rgb.endswith('CD7F32')


def test_missing_sheet_raises(tmp_path):
    path = tmp_path / 'empty.xlsx'
    pd.DataFrame({'a': [1]}).to_excel(path, sheet_name='other', index=False)
    with pytest.raises(ValueError, match='participation_ALL_clean'):
        process(str(path), str(tmp_path / 'out.xlsx'), log=lambda m: None)
